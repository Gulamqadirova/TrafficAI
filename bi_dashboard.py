from __future__ import annotations
from typing import Dict
import os
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.dates as mdates
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows
import config

CAR_C = "#ED7D31"; PED_C = "#2E75B6"
PALETTE = ["#2E75B6", "#ED7D31", "#548235", "#7030A0", "#C00000"]
plt.rcParams.update({"figure.dpi": 120, "font.size": 9, "axes.grid": True,
                     "grid.alpha": 0.3, "axes.spines.top": False,
                     "axes.spines.right": False})


def density_band(ratio: float) -> str:
    for lo, hi, lab in config.DENSITY_BANDS:
        if lo <= ratio < hi:
            return lab
    return "Congested"


def compute_kpis(df, cv, anomalies) -> Dict:
    df = df.copy(); df["timestamp"] = pd.to_datetime(df["timestamp"])
    peak_road = df.loc[df["road_occupancy"].idxmax()]
    site_cars = df.groupby("site")["car_count"].sum().sort_values(ascending=False)
    site_peds = df.groupby("site")["person_count"].sum().sort_values(ascending=False)
    cars_cv = cv["classes"].get("car", {})
    peds_cv = cv["classes"].get("person", {})
    return {
        "total_observations": int(len(df)),
        "zones_monitored": int(df["zone"].nunique()),
        "sites_monitored": int(df["site"].nunique()),
        "total_cars": int(df["car_count"].sum()),
        "total_people": int(df["person_count"].sum()),
        "peak_road_occupancy_pct": round(float(peak_road["road_occupancy"])*100, 1),
        "peak_zone": str(peak_road["zone"]),
        "peak_time": peak_road["timestamp"].strftime("%H:%M"),
        "peak_band": density_band(float(peak_road["road_occupancy"])),
        "anomalies_detected": int(len(anomalies)),
        "busiest_car_site": str(site_cars.index[0]),
        "busiest_ped_site": str(site_peds.index[0]),
        "cv_backend": cv.get("backend"),
        "cv_frames": cv.get("frames_processed"),
        "cv_car_crossings": cars_cv.get("total_crossings", 0),
        "cv_person_crossings": peds_cv.get("total_crossings", 0),
        "cv_car_conf": cars_cv.get("mean_confidence", 0.0),
        "cv_person_conf": peds_cv.get("mean_confidence", 0.0),
        "zone_peak_road": {z: round(v*100, 1) for z, v in
                           df.groupby("zone")["road_occupancy"].max().items()},
        "site_cars": {s: int(v) for s, v in site_cars.items()},
        "site_peds": {s: int(v) for s, v in site_peds.items()},
    }


# ---- charts ---------------------------------------------------------------- #
def chart_flow_trends(df, path, zone):
    d = df[df["zone"] == zone].copy(); d["timestamp"] = pd.to_datetime(d["timestamp"])
    fig, ax = plt.subplots(figsize=(8, 3.5))
    ax.plot(d["timestamp"], d["car_count"], color=CAR_C, label="Cars", linewidth=1.5)
    ax.plot(d["timestamp"], d["person_count"], color=PED_C, label="People", linewidth=1.5)
    ax.set_title(f"Car vs Pedestrian Flow - {zone}", fontweight="bold")
    ax.set_xlabel("Time of day"); ax.set_ylabel("Count per 5 min")
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%H:%M")); ax.legend(fontsize=8)
    fig.tight_layout(); fig.savefig(path); plt.close(fig); return path


def chart_road_heatmap(df, path):
    d = df.copy(); d["timestamp"] = pd.to_datetime(d["timestamp"]); d["hour"] = d["timestamp"].dt.hour
    piv = d.pivot_table(index="zone", columns="hour", values="road_occupancy", aggfunc="mean").fillna(0)
    fig, ax = plt.subplots(figsize=(8, 3.0))
    im = ax.imshow(piv.values, aspect="auto", cmap="YlOrRd", vmin=0, vmax=1.2)
    ax.set_xticks(range(0, 24, 2)); ax.set_xticklabels(range(0, 24, 2))
    ax.set_yticks(range(len(piv.index))); ax.set_yticklabels(piv.index, fontsize=7)
    ax.set_title("Mean Road Occupancy by Zone and Hour", fontweight="bold"); ax.set_xlabel("Hour")
    fig.colorbar(im, ax=ax, fraction=0.025).set_label("Road occupancy")
    fig.tight_layout(); fig.savefig(path); plt.close(fig); return path


def chart_anomaly_timeline(scored, zone, path):
    d = scored[scored["zone"] == zone].copy(); d["timestamp"] = pd.to_datetime(d["timestamp"])
    an = d[d["is_anomaly"]]
    fig, ax = plt.subplots(figsize=(8, 3.4))
    ax.plot(d["timestamp"], d["car_count"], color=CAR_C, linewidth=1.3, label="Cars")
    ax.scatter(an["timestamp"], an["car_count"], color="#C00000", s=28, zorder=5, label="Anomaly")
    ax.axhline(int(d["road_capacity"].iloc[0]), color="grey", ls="--", lw=1, label="Road capacity")
    ax.set_title(f"Traffic Anomaly Timeline - {zone}", fontweight="bold")
    ax.set_xlabel("Time of day"); ax.set_ylabel("Cars per 5 min")
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%H:%M")); ax.legend(fontsize=8)
    fig.tight_layout(); fig.savefig(path); plt.close(fig); return path


def chart_site_comparison(kpis, path):
    sites = list(kpis["site_cars"].keys())
    cars = [kpis["site_cars"][s] for s in sites]
    peds = [kpis["site_peds"].get(s, 0) for s in sites]
    x = np.arange(len(sites)); wbar = 0.38
    fig, ax = plt.subplots(figsize=(7.5, 3.3))
    ax.bar(x - wbar/2, cars, wbar, label="Cars", color=CAR_C)
    ax.bar(x + wbar/2, peds, wbar, label="People", color=PED_C)
    ax.set_xticks(x); ax.set_xticklabels(sites, rotation=15, fontsize=8)
    ax.set_title("Daily Volume by Site: Cars vs People", fontweight="bold")
    ax.set_ylabel("Cumulative count"); ax.legend(fontsize=8)
    fig.tight_layout(); fig.savefig(path); plt.close(fig); return path


def chart_forecast(history, fc, zone, path, cutoff=config.FORECAST_NOW):
    h = history[history["zone"] == zone].copy(); h["timestamp"] = pd.to_datetime(h["timestamp"])
    if cutoff: h = h[h["timestamp"] <= pd.Timestamp(cutoff)]
    fig, ax = plt.subplots(figsize=(8, 3.4))
    ax.plot(h["timestamp"], h["car_count"], color=CAR_C, lw=1.3, label="Cars (history)")
    ax.plot(fc["future_timestamps"], fc["car_forecast"], color="#9C3D00", lw=2,
            marker="o", ms=3, label="Car forecast")
    ax.axhline(0.85*fc["road_capacity"], color="#C00000", ls="--", lw=1, label="85% road capacity")
    ax.set_title(f"Next-Hour Traffic Forecast - {zone} (MAE={fc['car_mae']})", fontweight="bold")
    ax.set_xlabel("Time of day"); ax.set_ylabel("Cars per 5 min")
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%H:%M")); ax.legend(fontsize=8)
    fig.tight_layout(); fig.savefig(path); plt.close(fig); return path


def build_all_charts(df, scored, kpis, forecasts, out_dir=config.OUTPUT_DIR):
    busiest = max(kpis["zone_peak_road"], key=kpis["zone_peak_road"].get)
    return {
        "flow_trends": chart_flow_trends(df, os.path.join(out_dir, "chart_flow_trends.png"), busiest),
        "road_heatmap": chart_road_heatmap(df, os.path.join(out_dir, "chart_road_heatmap.png")),
        "anomaly_timeline": chart_anomaly_timeline(scored, busiest, os.path.join(out_dir, "chart_anomaly_timeline.png")),
        "site_comparison": chart_site_comparison(kpis, os.path.join(out_dir, "chart_site_comparison.png")),
        "forecast": chart_forecast(df, forecasts[busiest], busiest, os.path.join(out_dir, "chart_forecast.png")),
    }


# ---- Excel ----------------------------------------------------------------- #
def write_excel_report(df, scored, kpis, anomalies, forecasts, path=None):
    if path is None:
        path = os.path.join(config.OUTPUT_DIR, "Traffic_BI_Report.xlsx")
    wb = Workbook()
    hf = PatternFill("solid", fgColor="2E75B6"); hfont = Font(bold=True, color="FFFFFF")
    tfont = Font(bold=True, size=14, color="2E75B6")

    ws = wb.active; ws.title = "KPI Dashboard"
    ws["A1"] = "TPBI Platform - Executive KPI Dashboard"; ws["A1"].font = tfont
    ws["A2"] = "AI Traffic & Pedestrian Business Intelligence"; ws["A2"].font = Font(italic=True, color="808080")
    rows = [("Total observations", kpis["total_observations"]),
            ("Zones / sites", f"{kpis['zones_monitored']} / {kpis['sites_monitored']}"),
            ("Total cars (estimate)", kpis["total_cars"]),
            ("Total people (estimate)", kpis["total_people"]),
            ("Peak road occupancy (%)", kpis["peak_road_occupancy_pct"]),
            ("Peak zone / time", f"{kpis['peak_zone']} @ {kpis['peak_time']}"),
            ("Peak congestion band", kpis["peak_band"]),
            ("Anomalies detected", kpis["anomalies_detected"]),
            ("Busiest car site", kpis["busiest_car_site"]),
            ("Busiest pedestrian site", kpis["busiest_ped_site"]),
            ("CV backend / frames", f"{kpis['cv_backend']} / {kpis['cv_frames']}"),
            ("CV car crossings", kpis["cv_car_crossings"]),
            ("CV person crossings", kpis["cv_person_crossings"])]
    ws["A4"] = "Metric"; ws["B4"] = "Value"
    for c in ("A4", "B4"): ws[c].fill = hf; ws[c].font = hfont
    for i, (k, v) in enumerate(rows, 5): ws[f"A{i}"] = k; ws[f"B{i}"] = v
    ws.column_dimensions["A"].width = 30; ws.column_dimensions["B"].width = 28

    ws2 = wb.create_sheet("Hourly Flow")
    d = df.copy(); d["timestamp"] = pd.to_datetime(d["timestamp"]); d["hour"] = d["timestamp"].dt.hour
    hourly = d.groupby("hour")[["car_count", "person_count"]].mean().round(1).reset_index()
    for r in dataframe_to_rows(hourly, index=False, header=True): ws2.append(r)
    for c in ws2[1]: c.fill = hf; c.font = hfont; c.alignment = Alignment(horizontal="center")
    ch = LineChart(); ch.title = "Average Hourly Flow"; ch.height, ch.width = 9, 18
    ch.y_axis.title = "Avg count"; ch.x_axis.title = "Hour"
    data = Reference(ws2, min_col=2, max_col=3, min_row=1, max_row=hourly.shape[0]+1)
    cats = Reference(ws2, min_col=1, min_row=2, max_row=hourly.shape[0]+1)
    ch.add_data(data, titles_from_data=True); ch.set_categories(cats); ws2.add_chart(ch, "F2")

    ws3 = wb.create_sheet("Anomalies")
    if anomalies.empty: ws3["A1"] = "No anomalies detected."
    else:
        a = anomalies.copy(); a["timestamp"] = pd.to_datetime(a["timestamp"]).astype(str)
        for r in dataframe_to_rows(a, index=False, header=True): ws3.append(r)
        for c in ws3[1]: c.fill = hf; c.font = hfont
        ws3.column_dimensions["A"].width = 20; ws3.column_dimensions["H"].width = 32

    ws4 = wb.create_sheet("Forecast")
    ws4.append(["zone", "car_peak", "person_peak", "road_capacity",
                "congestion_predicted", "car_MAE", "person_MAE"])
    for c in ws4[1]: c.fill = hf; c.font = hfont
    for z, f in forecasts.items():
        ws4.append([z, f["car_predicted_peak"], f["person_predicted_peak"],
                    f["road_capacity"], f["congestion_predicted"],
                    f["car_mae"], f["person_mae"]])

    wb.save(path); return path


if __name__ == "__main__":
    import synthetic_data, video_analytics, anomaly_detection, predictive_analytics
    df = synthetic_data.generate_traffic_timeseries(); synthetic_data.generate_traffic_clip()
    cv = video_analytics.analyse_video(write_annotated=False)
    scored = anomaly_detection.detect_anomalies(df)
    anomalies = anomaly_detection.anomaly_summary(scored)
    fc = predictive_analytics.forecast_all_zones(df)
    k = compute_kpis(df, cv, anomalies)
    charts = build_all_charts(df, scored, k, fc)
    xlsx = write_excel_report(df, scored, k, anomalies, fc)
    print("KPIs sample:", {x: k[x] for x in list(k)[:9]})
    print("Charts:", list(charts.values())); print("Excel:", xlsx)
